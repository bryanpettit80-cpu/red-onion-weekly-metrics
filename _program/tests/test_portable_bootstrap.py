"""An existing legacy workspace can adopt portable state with empty intake."""
from pathlib import Path
import shutil

from openpyxl import load_workbook
import pytest

from red_onion_integrity import sha256_file
from red_onion_management import json_value, read_state
import red_onion_weekly_metrics as weekly
import red_onion_workflow as workflow
from test_portable_workflow import (
    arguments, make_workspace, output_hashes, published_workspace, store_for, table_cells,
)


def test_empty_intake_legacy_bootstrap_preserves_reports_history_and_management(
    published_workspace, tmp_path, monkeypatch,
):
    root = make_workspace(tmp_path / "Existing legacy workspace")
    history = root / "Archive/processed-daily-reports"
    history.mkdir()
    original = store_for(published_workspace).current()
    for day, item in original["days"].items():
        shutil.copy2(published_workspace / "Archive" / item["source"], history / f"Export {day}.xlsx")
    for path in (published_workspace / "Finished Reports").glob("*.xlsx"):
        shutil.copy2(path, root / "Finished Reports" / path.name)
    standard_master = root / "Finished Reports/Red_Onion_Server_Master.xlsx"
    master = standard_master.with_name("Weekly management review.xlsx")
    standard_master.rename(master)
    wb = load_workbook(master)
    try:
        ws, columns, rows = table_cells(wb, "ManagementTargets")
        target_row = next(row for row in rows if ws.cell(row, columns["Entity"]).value == "RC Richmond")
        ws.cell(target_row, columns["Weekly Sales Target"], 12345)
        wb.save(master)
    finally:
        wb.close()
    saved_management = json_value(read_state(master))
    reports_before = output_hashes(root)
    history_before = {p.name: sha256_file(p) for p in history.glob("*.xlsx")}

    def no_generation(*args, **kwargs):
        pytest.fail("Legacy bootstrap must not generate reports")
    with monkeypatch.context() as first_call:
        first_call.setattr(weekly, "write_public_workbook", no_generation)
        first_call.setattr(weekly, "write_master_workbook", no_generation)
        args = arguments(root)
        assert workflow.run(args) == []

    assert args._run_attempt_recorder.outcome == "history_imported"
    adopted = store_for(root).current()
    assert adopted["operation"] == "history_import"
    assert adopted["report_publication_id"] is None
    assert adopted["period"] is None
    assert adopted["outputs"] == []
    assert adopted["published_outputs"] == {}
    assert adopted["pending_report_refresh"] is True
    assert adopted["management_state"] == saved_management
    assert adopted["management_workbook_name"] == master.name
    assert len(adopted["days"]) == 6
    assert output_hashes(root) == reports_before
    assert {p.name: sha256_file(p) for p in history.glob("*.xlsx")} == history_before
    assert not list((root / "Daily Reports").iterdir())

    # The filename is observed management provenance, not a claim that the
    # portable workflow published this existing workbook. Missing it still
    # requires review; an ordinary call must not discard those saved inputs.
    saved_master = root / "Saved management backup.xlsx"
    master.rename(saved_master)
    try:
        missing = arguments(root)
        with pytest.raises(workflow.WorkflowAttention, match="management workbook is missing"):
            workflow.run(missing)
        assert missing._run_attempt_recorder.outcome == "needs_review"
        assert store_for(root).current() == adopted
        assert output_hashes(root) == {name: digest for name, digest in reports_before.items() if name != master.name}
    finally:
        saved_master.rename(master)

    # An ordinary second call refreshes accepted history without requiring a
    # new download. Reuse a valid master to keep this transition check focused.
    def master_copy(records, output_path, config, source_dir, public_start, public_end):
        assert json_value(config["_portable_management_state"]) == saved_management
        shutil.copy2(master, output_path)
        return output_path
    monkeypatch.setattr(weekly, "write_master_workbook", master_copy)
    refresh = arguments(root)
    assert len(workflow.run(refresh)) == 3
    assert refresh._run_attempt_recorder.outcome == "published"
    assert store_for(root).current()["report_publication_id"] is not None
    assert store_for(root).current()["pending_report_refresh"] is False
    assert {p.name: sha256_file(p) for p in history.glob("*.xlsx")} == history_before
    assert json_value(read_state(master)) == saved_management


def test_explicit_import_only_still_requires_a_source_folder(tmp_path):
    root = make_workspace(tmp_path / "Explicit import without source")
    args = arguments(root, migrate_history_only=True)
    with pytest.raises(ValueError, match="at least one --migrate-history-from"):
        workflow.run(args)
    assert store_for(root).current() is None
    assert output_hashes(root) == {}
