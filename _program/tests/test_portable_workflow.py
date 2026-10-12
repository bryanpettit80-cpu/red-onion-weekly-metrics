"""Exercise real source parsing, workbook generation and portable publication.

Only external storage failures and the operator's correction response are
simulated. A shared published fixture keeps the expensive generator runs low.
"""
from __future__ import annotations

from argparse import Namespace
from copy import copy
from datetime import date, time, timedelta
import json
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries
import pytest

from red_onion_integrity import sha256_file
from red_onion_management import read_state
from red_onion_transactions import ReportStore
import red_onion_weekly_metrics as weekly
import red_onion_workflow as workflow


WEEK_START = date(2026, 10, 6)
LOCATIONS = ("RC Richmond", "RC Virginia Beach")
PROGRAM = Path(__file__).resolve().parents[1]


def make_workspace(root: Path) -> Path:
    for name in ("Daily Reports", "Finished Reports", "Archive", "_program"):
        (root / name).mkdir(parents=True, exist_ok=True)
    shutil.copy2(PROGRAM / "red_onion_config.json", root / "_program/red_onion_config.json")
    return root


def arguments(root: Path, **overrides) -> Namespace:
    values = dict(
        config=str(root / "_program/red_onion_config.json"),
        input_dir=str(root / "Daily Reports"),
        output_dir=str(root / "Finished Reports"),
        archive_dir=str(root / "Archive"),
        week_start=None, week_end=None, rebuild_from_history=False,
        migrate_history_only=False, migrate_history_from=None,
        management_workbook=None, corrections_file=None, review_corrections=False,
    )
    return Namespace(**(values | overrides))


def store_for(root: Path) -> ReportStore:
    return ReportStore(root / "Archive", root / "Finished Reports", root / "Daily Reports")


def write_daily(
    root: Path, day: date, *, sales: float = 1500,
    locations: tuple[str, ...] = LOCATIONS, filename: str | None = None,
) -> Path:
    path = root / "Daily Reports" / (filename or f"Export {day.isoformat()}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Report(All)"
    ws.append([f"Date(s): {day:%m/%d/%Y}"])
    ws.append([
        "Store", "User", "Gross Sales", "Guest Count", "Check Average",
        "Wine Sales", "Rate of Sale by Guest Count", "Average Ticket Time", "Check Count",
    ])
    for location in locations:
        # One named person plus a matching restaurant total: totals must never
        # be added to the person's row to double the restaurant's business.
        for person in ("Alex Morgan", None):
            ws.append([location, person, sales, 150, sales / 150, 150, .02, time(0, 20), 40])
    wb.save(path)
    wb.close()
    return path


def write_week(root: Path, *, start: date = WEEK_START, locations=LOCATIONS) -> list[Path]:
    return [write_daily(root, start + timedelta(days=index), locations=locations) for index in range(6)]


def output_hashes(root: Path) -> dict[str, str]:
    return {p.name: sha256_file(p) for p in (root / "Finished Reports").glob("*.xlsx")}


def input_hashes(root: Path) -> dict[str, str]:
    return {p.name: sha256_file(p) for p in (root / "Daily Reports").glob("*.xlsx")}


def add_accepted_duplicate(root: Path, day: date = WEEK_START) -> Path:
    accepted = store_for(root).current()["days"][day.isoformat()]
    duplicate = root / "Daily Reports" / "Downloaded again with a different name.xlsx"
    shutil.copy2(root / "Archive" / accepted["source"], duplicate)
    return duplicate


def table_cells(wb, table_name: str):
    ws = next(ws for ws in wb.worksheets if table_name in ws.tables)
    left, top, right, bottom = range_boundaries(ws.tables[table_name].ref)
    headers = {str(ws.cell(top, col).value): col for col in range(left, right + 1)}
    return ws, headers, range(top + 1, bottom + 1)


def public_totals(root: Path, code: str, *, week_end: date = WEEK_START + timedelta(days=5)):
    wb = load_workbook(root / "Finished Reports" / f"Check_Wine_{code}{week_end:%m%d%y}.xlsx", data_only=True)
    try:
        ws = wb["Report(All)"]
        return tuple(ws.cell(3, col).value for col in range(3, 8))
    finally:
        wb.close()


@pytest.fixture(scope="session")
def published_workspace(tmp_path_factory) -> Path:
    root = make_workspace(tmp_path_factory.mktemp("portable-published"))
    inputs = write_week(root)
    args = arguments(root)
    outputs = workflow.run(args)
    assert args._run_attempt_recorder.outcome == "published"
    assert {p.name for p in outputs} == {
        "Check_Wine_RVA101126.xlsx", "Check_Wine_VB101126.xlsx", "Red_Onion_Server_Master.xlsx",
    }
    assert not any(path.exists() for path in inputs)
    return root


@pytest.fixture
def published_copy(published_workspace, tmp_path) -> Path:
    return Path(shutil.copytree(published_workspace, tmp_path / "Operating workspace"))


def test_real_six_day_source_generates_correct_complete_reports(published_workspace):
    root = published_workspace
    commit = store_for(root).current()
    assert commit["period"] == {"start": "2026-10-06", "end": "2026-10-11"}
    assert len(commit["days"]) == 6
    assert all(day["revision"] == 1 for day in commit["days"].values())
    for location, code in zip(LOCATIONS, ("RVA", "VB")):
        # Public report values are persisted numbers, not unevaluated formulas.
        assert public_totals(root, code) == pytest.approx((9000, 900, 10, 900, .10))
        records = [
            row for day in commit["days"].values() for row in day["records"]
            if row["location"] == location and row["is_location_total"]
        ]
        assert sum(row["check_count"] for row in records) == 240
    wb = load_workbook(root / "Finished Reports/Red_Onion_Server_Master.xlsx", data_only=True)
    try:
        ws, columns, rows = table_cells(wb, "WeeklyLocationMetrics")
        actual = {
            ws.cell(row, columns["Location"]).value: tuple(
                ws.cell(row, columns[label]).value
                for label in ("Gross Sales", "Guest Count", "Check Average", "Wine Sales", "Wine %")
            ) for row in rows
        }
        assert set(actual) == set(LOCATIONS)
        assert all(values == pytest.approx((9000, 900, 10, 900, .10)) for values in actual.values())
        snapshot = next(ws for ws in wb.worksheets if ws["A17"].value == "STORE SNAPSHOT")
        for row in (19, 20):
            assert tuple(snapshot.cell(row, col).value for col in range(5, 11)) == pytest.approx(
                (9000, 900, 240, 10, 37.5, 3.75)
            )
    finally:
        wb.close()


@pytest.mark.parametrize("missing_kind", ["day", "restaurant"])
def test_partial_week_does_not_publish_or_accept_inputs(tmp_path, missing_kind):
    root = make_workspace(tmp_path / missing_kind)
    if missing_kind == "day":
        for index in range(5):
            write_daily(root, WEEK_START + timedelta(days=index))
    else:
        write_week(root, locations=(LOCATIONS[0],))
    before = input_hashes(root)
    args = arguments(root)
    assert workflow.run(args) == []
    assert args._run_attempt_recorder.outcome == "waiting_for_inputs"
    assert store_for(root).current() is None
    assert input_hashes(root) == before
    assert output_hashes(root) == {}


def test_duplicate_input_is_archived_without_regenerating_reports(published_copy):
    root = published_copy
    previous = store_for(root).current()
    before = output_hashes(root)
    duplicate = add_accepted_duplicate(root)
    args = arguments(root)
    assert workflow.run(args) == []
    current = store_for(root).current()
    assert args._run_attempt_recorder.outcome == "already_processed"
    assert current["report_publication_id"] == previous["report_publication_id"]
    assert current["days"] == previous["days"]
    assert output_hashes(root) == before
    assert not duplicate.exists()


def test_incomplete_next_week_leaves_previous_publication_and_inputs_unchanged(published_copy):
    root = published_copy
    write_daily(root, WEEK_START + timedelta(days=7))
    previous = store_for(root).current()
    before, inputs = output_hashes(root), input_hashes(root)
    args = arguments(root)
    assert workflow.run(args) == []
    assert args._run_attempt_recorder.outcome == "waiting_for_inputs"
    assert store_for(root).current() == previous
    assert output_hashes(root) == before
    assert input_hashes(root) == inputs


def test_correction_requires_one_review_then_duplicate_retry_is_noop(published_copy, monkeypatch):
    root = published_copy
    prior = store_for(root).current()
    before = output_hashes(root)
    corrected = write_daily(root, WEEK_START, sales=1800, filename="Corrected Tuesday.xlsx")
    corrected_bytes = corrected.read_bytes()
    first = arguments(root)
    with pytest.raises(workflow.WorkflowAttention, match="correction choice"):
        workflow.run(first)
    assert first._run_attempt_recorder.outcome == "needs_review"
    assert store_for(root).current()["run_id"] == prior["run_id"]
    assert output_hashes(root) == before
    assert corrected.read_bytes() == corrected_bytes
    review = json.loads(Path(first._run_attempt_recorder.details["correction_review"]).read_text())
    item = review["corrections"][0]
    choice = next(index for index, option in enumerate(item["options"], 1) if option["file"] == corrected.name)
    prompts = []
    def choose():
        prompts.append(True)
        return str(choice)
    monkeypatch.setattr("builtins.input", choose)
    second = arguments(root, review_corrections=True)
    assert len(workflow.run(second)) == 3
    current = store_for(root).current()
    assert second._run_attempt_recorder.outcome == "published"
    assert len(prompts) == 1
    assert current["days"][WEEK_START.isoformat()]["revision"] == 2
    assert len(current["corrections"]) == 1
    assert public_totals(root, "RVA") == pytest.approx((9300, 900, 9300 / 900, 900, 900 / 9300))
    assert not corrected.exists()
    corrected.write_bytes(corrected_bytes)
    unchanged = output_hashes(root)
    third = arguments(root, review_corrections=True)
    assert workflow.run(third) == []
    repeated = store_for(root).current()
    assert third._run_attempt_recorder.outcome == "already_processed"
    assert repeated["report_publication_id"] == current["report_publication_id"]
    assert repeated["days"] == current["days"]
    assert output_hashes(root) == unchanged
    assert len(prompts) == 1


def test_renamed_resaved_master_edits_survive_duplicate_and_computer_roundtrip(published_copy, tmp_path):
    computer_a = published_copy
    store = store_for(computer_a)
    prior = store.current()
    master = computer_a / "Finished Reports/Red_Onion_Server_Master.xlsx"
    renamed = master.with_name("Weekly management review.xlsx")
    master.rename(renamed)
    wb = load_workbook(renamed)
    try:
        ws, columns, rows = table_cells(wb, "ManagementTargets")
        target_row = next(row for row in rows if ws.cell(row, columns["Entity"]).value == LOCATIONS[0])
        ws.cell(target_row, columns["Weekly Sales Target"], 10000)
        font = copy(ws.cell(target_row, columns["Entity"]).font)
        font.name = "Arial"
        ws.cell(target_row, columns["Entity"]).font = font
        ws, columns, rows = table_cells(wb, "OwnerRoster")
        first_row = next(iter(rows))
        ws.cell(first_row, columns["Owner Name"], "Taylor Manager")
        ws.cell(first_row, columns["Active"], "Yes")
        wb.save(renamed)
    finally:
        wb.close()
    add_accepted_duplicate(computer_a)
    args = arguments(computer_a)
    assert len(workflow.run(args)) == 3
    current = store_for(computer_a).current()
    assert args._run_attempt_recorder.outcome == "published"
    assert current["report_publication_id"] != prior["report_publication_id"]
    assert current["days"] == prior["days"]
    assert not master.exists()
    assert current["published_outputs"]["master"]["name"] == renamed.name
    management = read_state(renamed)
    assert management["targets"][LOCATIONS[0]]["gross_sales"] == 10000
    assert "Taylor Manager" in management["owners"]
    # Both hops copy only the operating workspace, without Git, an installation
    # directory, or an external trust anchor. Each PC resumes the saved receipt.
    computer_b = Path(shutil.copytree(computer_a, tmp_path / "Computer B"))
    add_accepted_duplicate(computer_b)
    args_b = arguments(computer_b)
    assert workflow.run(args_b) == []
    assert args_b._run_attempt_recorder.outcome == "already_processed"
    returned_a = Path(shutil.copytree(computer_b, tmp_path / "Computer A returned"))
    add_accepted_duplicate(returned_a)
    args_a = arguments(returned_a)
    assert workflow.run(args_a) == []
    assert args_a._run_attempt_recorder.outcome == "already_processed"
    final = store_for(returned_a).current()
    assert final["report_publication_id"] == current["report_publication_id"]
    assert final["days"] == current["days"]
    assert final["management_state"] == current["management_state"]
    assert output_hashes(returned_a) == output_hashes(computer_a)


def test_unavailable_cloud_input_preserves_previous_reports_and_receipt(published_copy, monkeypatch):
    root = published_copy
    incoming = write_daily(root, WEEK_START + timedelta(days=7))
    before, inputs = output_hashes(root), input_hashes(root)
    previous = store_for(root).current()
    original_capture = weekly.capture_regular_file
    def cloud_unavailable(path, *args, **kwargs):
        if Path(path) == incoming:
            raise OSError("Cloud file is unavailable while offline")
        return original_capture(path, *args, **kwargs)
    monkeypatch.setattr(weekly, "capture_regular_file", cloud_unavailable)
    args = arguments(root)
    with pytest.raises(OSError, match="Cloud file is unavailable"):
        workflow.run(args)
    assert args._run_attempt_recorder.outcome == "failed_before_publication"
    assert store_for(root).current() == previous
    assert output_hashes(root) == before
    assert input_hashes(root) == inputs


def test_commit_lock_failure_rolls_back_complete_generated_set(published_copy, monkeypatch):
    root = published_copy
    write_week(root, start=WEEK_START + timedelta(days=7))
    before, inputs = output_hashes(root), input_hashes(root)
    prior = store_for(root).current()
    def locked_receipt(self, run_id):
        raise PermissionError("Current reporting receipt is locked")
    monkeypatch.setattr(ReportStore, "_write_pointer", locked_receipt)
    args = arguments(root)
    with pytest.raises(PermissionError, match="receipt is locked"):
        workflow.run(args)
    assert args._run_attempt_recorder.outcome == "failed_before_publication"
    assert store_for(root).current() == prior
    assert output_hashes(root) == before
    assert input_hashes(root) == inputs
    assert not list((root / "Finished Reports").glob("*.previous"))


def test_explicit_full_week_location_closure_allows_complete_publication(tmp_path):
    root = make_workspace(tmp_path / "Declared closure")
    config_path = root / "_program/red_onion_config.json"
    config = json.loads(config_path.read_text())
    config["closed_dates"] = {LOCATIONS[1]: [(WEEK_START + timedelta(days=i)).isoformat() for i in range(6)]}
    config_path.write_text(json.dumps(config), encoding="utf-8")
    write_week(root, locations=(LOCATIONS[0],))
    args = arguments(root)
    assert len(workflow.run(args)) == 3
    assert args._run_attempt_recorder.outcome == "published"
    assert public_totals(root, "RVA") == pytest.approx((9000, 900, 10, 900, .10))
    closed = load_workbook(root / "Finished Reports/Check_Wine_VB101126.xlsx", data_only=True)
    try:
        assert "Declared closed" in closed["Report(All)"]["A1"].value
        assert closed["Report(All)"]["C3"].value is None
    finally:
        closed.close()
    assert len(store_for(root).current()["days"]) == 6
