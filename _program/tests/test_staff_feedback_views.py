from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries
import pytest

from red_onion_staff_feedback import build_staff_feedback
from red_onion_staff_views import (
    _money,
    write_staff_dashboard,
    write_staff_methodology,
    write_staff_scorecards,
    write_staff_weekly,
)


@pytest.mark.parametrize("value, signed, expected", [
    (100.625, False, "$100.63"),
    (100.625, True, "+$100.63"),
    (-100.625, True, "-$100.63"),
    (1.005, False, "$1.01"),
    (0.004, True, "$0.00"),
])
def test_text_money_matches_excel_half_up_currency_display(value, signed, expected):
    assert _money(value, signed=signed) == expected


def staff_model():
    weeks = [date(2026, 8, 2) + timedelta(weeks=offset) for offset in range(8)]
    source = []
    for week in weeks:
        for index in range(10):
            guests = 1 if index == 8 else 50
            spend = 85 if index in (7, 8, 9) else 115 if index == 6 else 100
            wine = .13 if index in (6, 8, 9) else .07 if index == 7 else .10
            source.append({"week_end": week, "location": "Example Store", "raw_user_name": f"Person {index}",
                           "display_name": f"Person {index}", "gross_sales": guests * spend,
                           "guest_count": guests, "wine_sales": guests * spend * wine,
                           "active_days": 1 if guests == 1 else 3, "feedback_peer_eligible": guests >= 25})
    config = {"locations": {"Example Store": {}}, "management_peer_reference": {"min_distinct_peers_per_week": 5}}
    return {"staff_feedback": build_staff_feedback(source, weeks, config), "methodology_version": "2026.10-v3"}, config


def test_all_printed_cards_fit_native_letter_with_unclipped_page_boundaries():
    model, config = staff_model()
    feedback = model["staff_feedback"]
    by_name = {row["display_name"]: row for row in feedback["rows"]}
    assert feedback["coaching_rows"]
    assert feedback["recognition_rows"]
    assert by_name["Person 6"]["overall_read"] == "Recognition opportunity"
    assert by_name["Person 7"]["overall_read"] == "Coaching conversation"
    assert by_name["Person 9"]["overall_read"] == "Mixed results"
    assert by_name["Person 9"]["coaching_prompt"]
    assert by_name["Person 9"]["recognition_prompt"]
    assert by_name["Person 8"]["data_coverage"] == "Early indication"
    assert by_name["Person 8"]["one_guest_weeks"] == 8
    assert by_name["Person 8"]["comparison_spg"] is not None
    assert by_name["Person 8"]["coaching_prompt"].startswith("Early indication:")
    assert by_name["Person 8"]["recognition_prompt"].startswith("Early indication:")
    assert all(row["compared_weeks"] == 8 for row in feedback["rows"])
    wb = Workbook()
    write_staff_scorecards(wb, model, config)
    write_staff_weekly(wb, model)
    write_staff_dashboard(wb, model, config, None)
    write_staff_methodology(wb, model, config)
    for name in ("Server Scorecards", "Weekly Performance", "Performance Dashboard", "Methodology"):
        ws = wb[name]
        area = str(ws.print_area).split("!", 1)[-1]
        min_column, start, max_column, finish = range_boundaries(area)
        repeated = [int(value.replace("$", "")) for value in ws.print_title_rows.split(":")]
        repeat_height = sum(ws.row_dimensions[row].height or 15 for row in range(repeated[0], repeated[1] + 1))
        end_rows = [brk.id for brk in ws.row_breaks.brk if start <= brk.id < finish] + [finish]
        max_height = 612 - 72 * (ws.page_margins.top + ws.page_margins.bottom)
        page_start = start
        for index, end in enumerate(end_rows):
            height = sum(ws.row_dimensions[row].height or 15 for row in range(page_start, end + 1))
            if index:
                height += repeat_height
            assert height <= max_height, (name, index + 1, height, max_height)
            page_start = end + 1
        assert all(cell.font.sz >= 9 for cells in ws.iter_rows(min_row=start, max_row=finish,
                   min_col=min_column, max_col=max_column) for cell in cells if cell.value not in (None, ""))
        assert ws.sheet_properties.pageSetUpPr.fitToPage is False
        if name in ("Server Scorecards", "Weekly Performance"):
            assert len(end_rows) == len(model["staff_feedback"]["rows"])


def test_empty_staff_data_still_has_readable_printable_explanation():
    wb = Workbook()
    model = {"staff_feedback": {"rows": [], "detail_rows": [], "week_ends": [],
                               "coaching_rows": [], "recognition_rows": [], "thresholds": {}}}
    write_staff_scorecards(wb, model, {})
    write_staff_weekly(wb, model)
    for name in ("Server Scorecards", "Weekly Performance"):
        ws = wb[name]
        assert any(cell.value == "No resolved named staff records are available." for cell in ws._cells.values())
        assert ws.print_area
