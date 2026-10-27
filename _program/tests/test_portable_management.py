from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
import pytest

from red_onion_management import find_master, read_state
import red_onion_weekly_metrics as weekly


def master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manager's chosen sheet name"
    targets = ["Entity", *(label for _, label in weekly.TARGET_FIELDS)]
    ws.append(targets)
    ws.append(["RC Richmond", 9000, 900, 10, .10, 6, 4])
    ws.add_table(Table(displayName="ManagementTargets", ref="A1:G2"))
    ws.append([])
    ws.append(["Owner Name", "Active"])
    ws.append(["Manager", "Yes"])
    ws.add_table(Table(displayName="OwnerRoster", ref="A4:B5"))
    ws.append([])
    headers = ["Action ID", "Entity Key", "Status", "Owner", "Due Date", "Context Notes", "Review Disposition", "Reviewed By", "Review Date"]
    ws.append(headers)
    ws.append(["action-1", "person|Richmond|Alex", "Review Needed", "Manager", date(2026, 10, 14), "Preserve my note", "Pending Review", None, None])
    ws.add_table(Table(displayName="ActionBoardTable", ref="A7:I8"))
    ws.append([])
    ws.append(headers)
    ws.append([None] * len(headers))
    ws.add_table(Table(displayName="ActionHistoryTable", ref="A10:I11"))
    wb.save(path)
    wb.close()


def test_management_inputs_survive_name_style_and_save_changes(tmp_path):
    path = tmp_path / "My management report.xlsx"
    master(path)
    previous = read_state(path)
    wb = load_workbook(path)
    ws = wb.active
    style = copy(ws["A1"].font)
    style.name = "Arial"
    style.sz = 16
    ws["A1"].font = style
    ws.column_dimensions["A"].width = 34
    ws["F8"] = "Updated manager note"
    ws["B8"] = "do not import altered generated identity"
    wb.save(path)
    wb.close()
    assert find_master(tmp_path) == path
    updated = read_state(path, previous)
    assert updated["active_actions"][0]["Context Notes"] == "Updated manager note"
    assert updated["active_actions"][0]["Entity Key"] == "person|Richmond|Alex"
    assert updated["targets"]["RC Richmond"]["gross_sales"] == 9000


def test_changed_or_removed_action_id_does_not_silently_lose_management_work(tmp_path):
    path = tmp_path / "Red_Onion_Server_Master.xlsx"
    master(path)
    prior = read_state(path)
    wb = load_workbook(path)
    wb.active["A8"] = "replacement-ID"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="Action IDs"):
        read_state(path, prior)


def test_multiple_different_master_copies_require_one_choice(tmp_path):
    first, second = tmp_path / "one.xlsx", tmp_path / "two.xlsx"
    master(first)
    master(second)
    wb = load_workbook(second)
    wb.active["F8"] = "Conflicting notes"
    wb.save(second)
    wb.close()
    with pytest.raises(ValueError, match="Several management"):
        find_master(tmp_path)
    assert find_master(tmp_path, str(second)) == second


def test_invalid_target_and_pasted_formula_are_not_accepted_as_blank(tmp_path):
    path = tmp_path / "report.xlsx"
    master(path)
    wb = load_workbook(path)
    wb.active["B2"] = "=1+1"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="not a formula"):
        read_state(path)


def test_unreadable_renamed_master_is_not_a_new_installation(tmp_path):
    renamed = tmp_path / "Weekly decisions.xlsx"
    renamed.write_bytes(b"unavailable workbook bytes")
    with pytest.raises(ValueError, match="unreadable"):
        find_master(tmp_path)
    with pytest.raises(ValueError, match="Cannot read management"):
        find_master(tmp_path, expected_name=renamed.name)


@pytest.mark.parametrize("title", [
    "ACTION HISTORY | 0 completed or dismissed item(s)",
    "ACTION HISTORY | 0 completed; 0 dismissed; 0 cleared without completed follow-up",
])
def test_empty_history_accepts_legacy_and_current_labels_without_dropping_rows(title):
    import red_onion_weekly_metrics as weekly
    from red_onion_management import _actions
    wb = Workbook()
    ws = wb.active
    ws["C3"] = title
    for column, label in enumerate(weekly.ACTION_HEADERS, 1):
        ws.cell(4, column, label)
    assert _actions(wb, "ActionHistoryTable", {"Action ID"}) == []
    ws["A5"] = "preserved-manager-action"
    with pytest.raises(ValueError, match="valid empty section"):
        _actions(wb, "ActionHistoryTable", {"Action ID"})
    wb.close()
