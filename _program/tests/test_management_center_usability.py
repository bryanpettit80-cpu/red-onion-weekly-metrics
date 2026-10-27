from __future__ import annotations

from copy import deepcopy
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries

import red_onion_weekly_metrics as metrics


def center_for(actions, history=()):
    wb = Workbook()
    end = date(2026, 10, 25)
    days = tuple(end - timedelta(days=offset) for offset in range(6))
    readiness = metrics.LatestWeekReadiness(
        latest_week_end=end,
        latest_location_rows=(),
        configured_locations=("RC Richmond",),
        location_gaps=(),
        expected_dates=days,
        received_dates=frozenset(days),
        missing_dates=(),
        ready=True,
    )
    metrics.write_management_center_sheet(
        wb, {}, [], list(actions), list(history),
        {"locations": {"RC Richmond": {"short_code": "RVA"}}}, readiness, [],
    )
    return wb["Management Center"]


def action(action_id="review-1", **changes):
    item = {
        "Action ID": action_id, "Status": "Review Needed", "Priority": "Review",
        "Location": "RC Richmond", "Person / Area": "Taylor Guest",
        "Action": "Context Review", "Why It Matters": "Sales per guest changed.",
        "Recommended Next Step": "Review recent shifts with the server.",
        "Review Disposition": "Pending Review", "Signal State": "Current",
    }
    item.update(changes)
    return item


def test_history_distinguishes_cleared_signals_from_completed_follow_up():
    history = [
        action("a", **{"Signal State": "Cleared", "Status": "Open"}),
        action("b", **{"Signal State": "Cleared"}),
        action("c", **{"Signal State": "Cleared", "Status": "Complete"}),
        action("d", **{"Signal State": "Cleared", "Status": "Dismissed"}),
    ]
    before = deepcopy(history)
    ws = center_for([], history)
    heading = next(cell.value for cell in ws["C"] if str(cell.value).startswith("ACTION HISTORY |"))
    assert "1 completed; 1 dismissed; 2 cleared without completed follow-up" in heading
    assert history == before
    _, header, _, last = range_boundaries(ws.tables["ActionHistoryTable"].ref)
    assert [ws.cell(row, 4).value for row in range(header + 1, last + 1)] == [
        "Open", "Review Needed", "Complete", "Dismissed",
    ]


def test_assignment_indicators_remain_linked_to_editable_actions():
    ws = center_for([action(), action("assigned", Owner="Manager", **{"Due Date": date(2026, 10, 28)})])
    _, header, _, last = range_boundaries(ws.tables["ActionBoardTable"].ref)
    for address, column in (("L6", "E"), ("L7", "F")):
        formula = ws[address].value
        assert f'${column}${header + 1}:${column}${last},""' in formula
        assert '"<>Complete"' in formula and '"<>Dismissed"' in formula
    assert ws.cell(header + 1, 5).value is None
    assert ws.cell(header + 1, 6).value is None
    assert ws.cell(header + 1, 5).protection.locked is False
    assert center_for([])["L6"].value == 0
    assert center_for([])["L7"].value == 0


def test_print_handout_has_independent_width_and_keeps_action_cards_together():
    ws = center_for([action(), action("second")], [action("old", **{"Signal State": "Cleared"})])
    _, _, _, history_last = range_boundaries(ws.tables["ActionHistoryTable"].ref)
    print_ref = str(ws.print_area).split("!")[1].replace("$", "")
    left, first, right, last = range_boundaries(print_ref)
    assert (left, right) == (25, 32)
    assert first > history_last
    assert ws.page_setup.fitToHeight == 0
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_margins.top == ws.page_margins.bottom == 0.4
    assert ws.print_title_rows == f"${first}:${first + 3}"
    assert len(ws.row_breaks.brk) == 1
    assert ws.row_breaks.brk[0].id == first + 12
    assert ws["K9"].hyperlink.target.endswith(f"!Y{first}")
    assert ws.cell(first + 3, 25).hyperlink.target.startswith("#'Management Center'!C")
    for row in ws.iter_rows(min_row=first, max_row=last, min_col=25, max_col=32):
        for cell in row:
            if cell.value is not None:
                assert cell.font.sz == 11
                assert cell.protection.locked is True
    assert ws.cell(first + 7, 25).data_type == "f"
    assert "Unassigned" in ws.cell(first + 7, 25).value


def test_guide_print_settings_survive_finalizer_without_screen_changes():
    wb = Workbook()
    metrics.write_how_to_use_sheet(wb)
    ws = wb["How to Use"]
    screen_dimensions = (
        {key: dim.width for key, dim in ws.column_dimensions.items()},
        {key: dim.height for key, dim in ws.row_dimensions.items()},
        set(str(merged) for merged in ws.merged_cells.ranges),
    )
    metrics.configure_management_print_layout(ws)
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.scale is None
    assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_LETTER)
    assert ws.print_title_rows == "$1:$1"
    assert [item.id for item in ws.row_breaks.brk] == [12, 26, 36, 43, 52]
    assert ws["A4"].font.sz == ws["A9"].font.sz == 11
    assert screen_dimensions == (
        {key: dim.width for key, dim in ws.column_dimensions.items()},
        {key: dim.height for key, dim in ws.row_dimensions.items()},
        set(str(merged) for merged in ws.merged_cells.ranges),
    )
