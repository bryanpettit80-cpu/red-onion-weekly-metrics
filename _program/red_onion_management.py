"""Portable management input, independent of workbook formatting and file names.

Generated metrics are never imported as new reporting facts. After the first
legacy import, only the documented editable fields are merged into the saved
management snapshot by Action ID. Source-based generation replaces everything
else. The snapshot travels with the operating workspace.
"""
from __future__ import annotations

import copy
from datetime import date, datetime
import json
import math
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


EDITABLE_ACTION_FIELDS = (
    "Status", "Owner", "Due Date", "Context Notes", "Review Disposition",
    "Reviewed By", "Review Date",
)
REQUIRED_TABLES = {"ManagementTargets", "OwnerRoster"}


def json_value(value: Any) -> Any:
    """Use portable ISO dates; weekly.as_date accepts these on restoration."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Management state contains a non-finite number.")
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ValueError(f"Unsupported management value: {type(value).__name__}")


def empty_state() -> dict[str, Any]:
    import red_onion_weekly_metrics as weekly
    return {
        "targets": {}, "owners": [], "owner_roster": [],
        "owner_roster_capacity": weekly.OWNER_ROSTER_MIN_EDIT_ROWS,
        "active_actions": [], "action_history": [], "evidence_by_action_id": {},
    }


def table_names(path: Path) -> set[str]:
    """Identify the workbook by its tables, not its path or presentation."""
    with zipfile.ZipFile(path) as archive:
        return {
            str(ElementTree.fromstring(archive.read(name)).get("name"))
            for name in archive.namelist()
            if name.startswith("xl/tables/") and name.endswith(".xml")
        }


def find_master(output_dir: Path, selected: str | None = None, *, expected_name: str | None = None) -> Path | None:
    if selected:
        chosen = Path(selected).expanduser().resolve()
        if chosen.parent != output_dir.resolve() or chosen.suffix.lower() != ".xlsx":
            raise ValueError("The selected management workbook must be in Finished Reports.")
        if not REQUIRED_TABLES.issubset(table_names(chosen)):
            raise ValueError("The selected workbook is missing its management input tables.")
        return chosen
    candidates: list[Path] = []
    unreadable: list[Path] = []
    for path in sorted(output_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            names = table_names(path)
        except (OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            # Do not mistake an unavailable or damaged master for a new installation.
            if "master" in path.stem.casefold() or path.name == expected_name:
                raise ValueError(f"Cannot read management workbook {path.name}: {exc}") from exc
            unreadable.append(path)
            continue
        if REQUIRED_TABLES.issubset(names):
            candidates.append(path)
        elif "master" in path.stem.casefold() or path.name == expected_name:
            raise ValueError(
                f"{path.name} is missing management input tables. Restore them or "
                "select the intact workbook with --management-workbook."
            )
    if len(candidates) > 1:
        states = [json.dumps(json_value(read_state(p)), sort_keys=True) for p in candidates]
        if len(set(states)) != 1:
            raise ValueError(
                "Several management workbooks contain different saved inputs: "
                + ", ".join(p.name for p in candidates)
                + ". Choose one with --management-workbook; no edits were discarded."
            )
        standard = output_dir / "Red_Onion_Server_Master.xlsx"
        return standard if standard in candidates else candidates[0]
    if not candidates and unreadable:
        raise ValueError("Cannot identify the management workbook while these workbooks are unreadable: " + ", ".join(p.name for p in unreadable))
    return candidates[0] if candidates else None


def _table(wb: Any, name: str, required: set[str]) -> tuple[list[dict[str, Any]], int]:
    matches = [ws for ws in wb.worksheets if name in ws.tables]
    if len(matches) != 1:
        raise ValueError(f"Management workbook must contain exactly one {name} table.")
    ws = matches[0]
    left, top, right, bottom = range_boundaries(ws.tables[name].ref)
    headers = [str(ws.cell(top, col).value or "").strip() for col in range(left, right + 1)]
    if len(headers) != len(set(headers)) or not required.issubset(headers):
        raise ValueError(f"{name} has missing or ambiguous column names.")
    rows: list[dict[str, Any]] = []
    for row in range(top + 1, bottom + 1):
        item: dict[str, Any] = {}
        for col, header in enumerate(headers, left):
            cell = ws.cell(row, col)
            if cell.data_type == "f" and header in required:
                raise ValueError(f"Enter a value, not a formula, in {name}.{header} row {row}.")
            item[header] = cell.value
        if any(value is not None and str(value).strip() for value in item.values()):
            rows.append(item)
    return rows, bottom - top


def _action_ids(rows: list[dict[str, Any]], label: str) -> dict[str, dict[str, Any]]:
    keyed: dict[str, dict[str, Any]] = {}
    for row in rows:
        key = str(row.get("Action ID") or "").strip()
        if not key or key in keyed:
            raise ValueError(f"{label} contains a missing or duplicate Action ID.")
        keyed[key] = row
    return keyed


def _actions(wb: Any, name: str, required: set[str]) -> list[dict[str, Any]]:
    if any(name in ws.tables for ws in wb.worksheets):
        return _table(wb, name, required)[0]
    # Pre-portability workbooks omit the table when the queue is empty.
    # Accept only the recognizable, actually blank empty section, never a
    # missing table that still has saved manager rows underneath its header.
    import red_onion_weekly_metrics as weekly
    titles, offset = (
        (("CURRENT ACTIONS | 0 open review item(s)",), 2)
        if name == "ActionBoardTable" else ((
            "ACTION HISTORY | 0 completed or dismissed item(s)",
            "ACTION HISTORY | 0 completed; 0 dismissed; 0 cleared without completed follow-up",
        ), 1)
    )
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_col=3, max_col=3):
            cell = row[0]
            if cell.value not in titles:
                continue
            header = cell.row + offset
            names = [ws.cell(header, col).value for col in range(1, len(weekly.ACTION_HEADERS) + 1)]
            blank = all(ws.cell(header + 1, col).value is None for col in range(1, len(weekly.ACTION_HEADERS) + 1))
            if names == list(weekly.ACTION_HEADERS) and blank:
                return []
    raise ValueError(f"Management workbook is missing its {name} table or valid empty section.")


def read_state(path: Path, previous: dict[str, Any] | None = None) -> dict[str, Any]:
    """Read required business fields while allowing saves, styles and table moves."""
    import red_onion_weekly_metrics as weekly
    wb = load_workbook(path, data_only=False)
    try:
        state = empty_state()
        labels = {label for _, label in weekly.TARGET_FIELDS}
        targets, _ = _table(wb, "ManagementTargets", {"Entity", *labels})
        for row in targets:
            entity = str(row.get("Entity") or "").strip()
            if not entity or entity in state["targets"]:
                raise ValueError("Management targets contain a missing or duplicate Entity.")
            values: dict[str, float | None] = {}
            for field, label in weekly.TARGET_FIELDS:
                value = row.get(label)
                if value is None or value == "":
                    values[field] = None
                elif isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                    raise ValueError(f"{entity}: {label} must be a nonnegative number or blank.")
                else:
                    values[field] = float(value) * (60 if field == "average_ticket_time_seconds" else 1)
            state["targets"][entity] = values
        roster, capacity = _table(wb, "OwnerRoster", {"Owner Name", "Active"})
        state["owner_roster"] = weekly.normalize_owner_roster(roster)
        if not 1 <= capacity <= weekly.OWNER_ROSTER_MAX_ROWS:
            raise ValueError("Owner Roster has an unsupported number of input rows.")
        state["owners"] = weekly.active_owner_names(state["owner_roster"])
        state["owner_roster_capacity"] = capacity
        active = _actions(wb, "ActionBoardTable", {"Action ID", *EDITABLE_ACTION_FIELDS})
        history = _actions(wb, "ActionHistoryTable", {"Action ID"})
        active_ids = _action_ids(active, "Current actions")
        _action_ids(history, "Action history")
        if previous is not None:
            previous_ids = _action_ids(previous.get("active_actions", []), "Saved current actions")
            if active_ids.keys() != previous_ids.keys():
                raise ValueError(
                    "Management Action IDs changed or rows were removed. Restore the current "
                    "action rows before reporting; saved management work is preserved."
                )
            state["active_actions"] = [
                {**copy.deepcopy(previous_ids[key]), **{field: row.get(field) for field in EDITABLE_ACTION_FIELDS}}
                for key, row in active_ids.items()
            ]
            state["action_history"] = copy.deepcopy(previous.get("action_history", []))
            state["evidence_by_action_id"] = copy.deepcopy(previous.get("evidence_by_action_id", {}))
        else:
            state["active_actions"] = active
            state["action_history"] = history
            if "Evidence Detail" in wb.sheetnames:
                evidence = weekly.records_from_sheet(wb["Evidence Detail"], "Evidence ID")
                state["evidence_by_action_id"] = {
                    str(row["Action ID"]): row for row in evidence if row.get("Action ID")
                }
        state["active_actions"] = weekly.validate_action_board_records(
            state["active_actions"], allowed_reviewers=state["owners"]
        )
        state["action_history"] = weekly.validate_action_board_records(state["action_history"])
        return state
    finally:
        wb.close()
