"""Plain-language views of the descriptive staff feedback model.

Screen tables retain every resolved person. Compact print views use the same
model on the same sheets, so staff discussions do not require tiny scaled grids.
The application owns navigation, protection, identity and metric calculation.
"""

from __future__ import annotations

from datetime import date, datetime
from textwrap import wrap
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


BURGUNDY = "7A1E1E"
INK = "262626"
MUTED = "595959"
BLUE = "E8F0FA"
GREEN = "E8F3E8"
AMBER = "FFF4D6"
PALE = "F5F5F5"
CURRENCY = '$#,##0.00;[Red]-$#,##0.00;"$0.00"'
SIGNED_CURRENCY = '+$0.00;[Red]-$0.00;"$0.00"'
PRINT_START_COLUMN = 25
PRINT_END_COLUMN = 32


def _safe(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _text(ws, row: int, first: int, last: int, value: Any, *, fill: str | None = None,
          color: str = INK, bold: bool = False, size: int = 10,
          align: str = "left", height: float | None = None):
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    cell = ws.cell(row=row, column=first, value=_safe(value))
    cell.font = Font(name="Calibri", size=size, color=color, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if height is not None:
        ws.row_dimensions[row].height = height
    return cell


def _sheet(wb, name: str, title: str, end: int = 12):
    if name in wb:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    _text(ws, 1, 1, max(12, end), title, fill=BURGUNDY, color="FFFFFF", bold=True,
          size=16, height=32)
    ws.row_dimensions[2].height = 24  # reserved for the application's workbook menu
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = "A5"
    ws.sheet_view.topLeftCell = "A1"
    return ws


def _letter_print(ws, area: str, repeat: str):
    ws.print_area = area
    ws.print_title_rows = repeat
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.scale = 100
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    # Native font sizes are intentional. Fit-to-page must not silently shrink
    # wide screen tables or already paginated cards when the file is finalized.
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.page_margins.header = ws.page_margins.footer = 0.1
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.right.text = ws.title
    ws.oddFooter.right.size = 9


def _table(ws, name: str, header: int, headers: list[str], records: list[list[Any]],
           *, row_height: int = 56):
    for column, label in enumerate(headers, 1):
        _text(ws, header, column, column, label, fill=BURGUNDY, color="FFFFFF",
              bold=True, align="center")
    ws.row_dimensions[header].height = 42
    for row_index, values in enumerate(records, header + 1):
        for column, value in enumerate(values, 1):
            _text(ws, row_index, column, column, value,
                  fill="FFFFFF" if (row_index - header) % 2 else PALE)
        ws.row_dimensions[row_index].height = row_height
    if records:
        table = Table(displayName=name,
                      ref=f"A{header}:{get_column_letter(len(headers))}{header + len(records)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def _feedback(model: dict[str, Any]) -> dict[str, Any]:
    return model.get("staff_feedback") or {"rows": [], "detail_rows": [], "week_ends": [],
                                           "coaching_rows": [], "recognition_rows": [],
                                           "checkin_rows": [], "thresholds": {}}


def _key(item: dict[str, Any]) -> tuple[str, str]:
    return str(item["location"]), str(item["raw_user_name"])


def _date(value: Any) -> str:
    return value.strftime("%m/%d/%Y") if isinstance(value, (date, datetime)) else "Not recorded"


def _money(value: float | None, *, signed: bool = False) -> str:
    if value is None:
        return "Unavailable"
    if signed:
        return f"{'+' if value > 0 else '-' if value < 0 else ''}${abs(value):,.2f}"
    return f"${value:,.2f}"


def _wine(value: float | None) -> float | None:
    return value * 100 if value is not None else None


def _guests(value: Any) -> str:
    return f"{float(value or 0):,.0f}"


def _activity(item: dict[str, Any]) -> str:
    return ("Recorded activity in latest week" if item.get("present_latest_week")
            else "No recorded activity latest week")


def _sample(item: dict[str, Any]) -> str:
    return (f"{item.get('compared_weeks', 0)} compared weeks | "
            f"{_guests(item.get('compared_guests'))} guests | "
            f"{item.get('data_coverage', 'No comparison')}")


def _evidence(item: dict[str, Any]) -> str:
    return (f"Spend: {_money(item.get('comparison_spg'))} vs {_money(item.get('peer_spg'))} "
            f"({_money(item.get('average_spg_gap'), signed=True)})\n"
            f"Wine / $100: {_money(_wine(item.get('comparison_wine_pct')))} vs "
            f"{_money(_wine(item.get('peer_wine_pct')))} "
            f"({_money(_wine(item.get('average_wine_gap')), signed=True)})\n"
            f"Steadiness: {item.get('consistency', 'Unavailable')}")


def _prompt(item: dict[str, Any], kind: str) -> str:
    value = item.get("feedback_text") if kind == "historical" else item.get(f"{kind}_prompt")
    return str(value or item.get("checkin_prompt") or
               "Ask about the shifts behind these results and agree on one practical next step.")


def _line_height(text: str, chars: int, *, minimum: int = 24) -> int:
    lines = sum(max(1, len(wrap(line, max(10, chars), break_long_words=True)))
                for line in str(text).splitlines())
    return max(minimum, lines * 12 + 10)


def _print_header(ws, start: int, title: str, note: str) -> None:
    for column in range(PRINT_START_COLUMN, PRINT_END_COLUMN + 1):
        ws.column_dimensions[get_column_letter(column)].width = 12
    _text(ws, start, 25, 32, title, fill=BURGUNDY, color="FFFFFF", bold=True,
          size=14, height=28)
    _text(ws, start + 1, 25, 32, note, color=MUTED, height=28)
    ws.row_dimensions[start + 2].height = 8


def _scorecard_print(ws, rows: list[dict[str, Any]], start: int) -> int:
    _print_header(ws, start, "Staff conversation scorecards",
                  "One person per page. Actual and typical teammate use the same weeks. All amounts are dollars.")
    current = start + 3
    if not rows:
        _text(ws, current, 25, 32, "No resolved named staff records are available.", height=36)
        return current
    for index, item in enumerate(rows):
        if index:
            ws.row_breaks.append(Break(id=current - 1))
        _text(ws, current, 25, 32, f"{item['display_name']} | {item['location']}",
              fill=BLUE, bold=True, size=14, height=30)
        _text(ws, current + 1, 25, 32,
              f"{_activity(item)} | Last worked: {_date(item.get('last_worked_date'))}",
              color=MUTED, height=24)
        sample_note = str(item.get("coverage_note") or _sample(item))
        _text(ws, current + 2, 25, 32, sample_note, fill=PALE,
              height=_line_height(sample_note, 100, minimum=38))
        for first, last, value in ((25, 27, "Selling result"), (28, 28, "Actual"),
                                   (29, 30, "Typical teammate"), (31, 32, "Difference")):
            _text(ws, current + 3, first, last, value, fill=BURGUNDY, color="FFFFFF",
                  bold=True, align="center", height=30)
        metrics = (("Spend per guest", "comparison_spg", "peer_spg", "average_spg_gap", False),
                   ("Wine sales per $100", "comparison_wine_pct", "peer_wine_pct", "average_wine_gap", True))
        for offset, (label, actual, peer, gap, wine) in enumerate(metrics, 4):
            _text(ws, current + offset, 25, 27, label, fill=PALE, bold=True, height=30)
            for first, last, field in ((28, 28, actual), (29, 30, peer), (31, 32, gap)):
                value = _wine(item.get(field)) if wine else item.get(field)
                cell = _text(ws, current + offset, first, last,
                             value if value is not None else "Unavailable", align="center", fill=PALE)
                cell.number_format = SIGNED_CURRENCY if field == gap else CURRENCY
        _text(ws, current + 6, 25, 32, f"Conversation guide: {item.get('overall_read', 'Routine check-in')}",
              fill=BLUE, bold=True, height=28)
        row = current + 7
        prompts = [("Coaching", item.get("coaching_prompt")),
                   ("Recognition", item.get("recognition_prompt"))]
        prompts = [(label, text) for label, text in prompts if text]
        if item.get("near_cutoff") or not prompts:
            prompts.append(("Check-in", _prompt(item, "checkin")))
        for label, text in prompts:
            _text(ws, row, 25, 32, f"{label}: {text}", fill=GREEN if label == "Recognition" else PALE,
                  height=_line_height(f"{label}: {text}", 100, minimum=40))
            row += 1
        _text(ws, row, 25, 32,
              f"Week-to-week steadiness: {item.get('consistency_note') or item.get('consistency', 'Unavailable')}",
              height=42)
        _text(ws, row + 1, 25, 32,
              f"All recorded work: {item.get('observed_weeks', 0)} weeks / "
              f"{_guests(item.get('observed_guests'))} guests; spend {_money(item.get('observed_spg'))}; "
              f"wine {_money(_wine(item.get('observed_wine_pct')))} per $100. "
              "This total can include weeks without a valid teammate comparison.", color=MUTED, height=40)
        _text(ws, row + 2, 25, 32,
              "Ask about section, shift, party size, promotions and guest preferences before interpreting a gap. "
              "Agree on one practical step and revisit it after more recorded shifts. "
              "Near-threshold results are a check-in, not a firm positive or negative conclusion.",
              fill=AMBER, height=48)
        current = row + 4
    return current - 1


def write_staff_scorecards(wb, model: dict[str, Any], config: dict[str, Any]) -> None:
    feedback = _feedback(model)
    rows = list(feedback["rows"])
    ws = _sheet(wb, "Server Scorecards", "Staff scorecards | clear results and useful conversations", 18)
    _text(ws, 3, 1, 18, "Every resolved named server is included, including small samples and people "
          "without recorded activity in the latest week. Filter Location or Server to find a person.",
          fill=PALE, height=32)
    _text(ws, 4, 1, 18, "Actual and typical teammate use the same compared weeks. Difference = actual minus "
          "teammate. Wine sales per $100 means dollars of wine in each $100 of sales. Early indications "
          "start a conversation; they are not firm conclusions. Use Print for one readable page per person.",
          fill=BLUE, height=44)
    headers = ["Location", "Server", "Latest week activity", "Data coverage", "Compared weeks", "Compared guests",
               "Spend per guest", "Typical teammate spend", "Spend difference", "Wine sales per $100",
               "Typical teammate wine", "Wine difference per $100", "Week-to-week steadiness", "Conversation guide",
               "Suggested conversation", "Last worked", "Sample context", "Steadiness context"]
    records = []
    for item in rows:
        records.append([item["location"], item["display_name"], _activity(item), item.get("data_coverage"),
                        item.get("compared_weeks", 0), item.get("compared_guests", 0), item.get("comparison_spg"),
                        item.get("peer_spg"), item.get("average_spg_gap"), _wine(item.get("comparison_wine_pct")),
                        _wine(item.get("peer_wine_pct")), _wine(item.get("average_wine_gap")), item.get("consistency"),
                        item.get("overall_read"), item.get("feedback_text") or _prompt(item, "checkin"),
                        item.get("last_worked_date"), item.get("coverage_note"), item.get("consistency_note")])
    _table(ws, "ServerScorecardsTable", 6, headers, records, row_height=84)
    widths = [21, 25, 30, 19, 12, 14, 16, 17, 16, 18, 18, 18, 22, 30, 72, 15, 64, 64]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for row_index in range(7, 7 + len(rows)):
        for column in range(7, 13):
            ws.cell(row_index, column).number_format = SIGNED_CURRENCY if column in (9, 12) else CURRENCY
        ws.cell(row_index, 6).number_format = "#,##0"
        ws.cell(row_index, 16).number_format = "mm/dd/yyyy"
    detail_start = len(rows) + 10
    _text(ws, detail_start - 1, 1, 10, "ALL RECORDED WORK | totals before a valid teammate comparison is required",
          fill=BLUE, bold=True, height=28)
    detail_headers = ["Location", "Server", "Recorded weeks", "Recorded guests", "Spend per guest (all work)",
                      "Wine per $100 (all work)", "Compared weeks", "Compared guests", "Last worked", "Source names"]
    details = [[item["location"], item["display_name"], item.get("observed_weeks", 0), item.get("observed_guests", 0),
                item.get("observed_spg"), _wine(item.get("observed_wine_pct")), item.get("compared_weeks", 0),
                item.get("compared_guests", 0), item.get("last_worked_date"),
                ", ".join(item.get("source_raw_names") or [item["raw_user_name"]])] for item in rows]
    _table(ws, "ServerScorecardDetailTable", detail_start, detail_headers, details, row_height=44)
    for row_index in range(detail_start + 1, detail_start + len(rows) + 1):
        for column in (5, 6):
            ws.cell(row_index, column).number_format = CURRENCY
        ws.cell(row_index, 9).number_format = "mm/dd/yyyy"
        ws.row_dimensions[row_index].height = _line_height(
            str(ws.cell(row_index, 10).value or ""), 17, minimum=44)
    print_start = detail_start + len(rows) + 4
    print_end = _scorecard_print(ws, rows, print_start)
    link = _text(ws, 5, 1, 18, "Open the printable individual scorecards →", color=BURGUNDY, bold=True, height=24)
    link.hyperlink = f"#'Server Scorecards'!Y{print_start}"
    ws.freeze_panes = "C7"
    _letter_print(ws, f"Y{print_start}:AF{print_end}", f"{print_start}:{print_start + 1}")


def _weekly_cell(detail: dict[str, Any] | None) -> str:
    if detail is None:
        return "No recorded activity"
    result = (f"Spend {_money(detail.get('check_average'))} "
              f"({_money(detail.get('spg_gap'), signed=True)})\n"
              f"Wine {_money(_wine(detail.get('wine_pct')))} "
              f"({_money(_wine(detail.get('wine_gap')), signed=True)})\n"
              f"{_guests(detail.get('guest_count'))} guests")
    return result


def _weekly_print(ws, rows: list[dict[str, Any]], detail_map: dict, weeks: list,
                  start: int) -> int:
    _print_header(ws, start, "Weekly staff results",
                  "Spend = sales per guest. Wine = wine sales per $100. Difference = actual minus typical teammate.")
    current = start + 3
    if not rows:
        _text(ws, current, 25, 32, "No resolved named staff records are available.", height=36)
        return current
    for index, item in enumerate(rows):
        if index:
            ws.row_breaks.append(Break(id=current - 1))
        _text(ws, current, 25, 32, f"{item['display_name']} | {item['location']}",
              fill=BLUE, bold=True, size=14, height=28)
        _text(ws, current + 1, 25, 32, f"{_activity(item)}. {_sample(item)}.", height=24)
        ws.row_dimensions[current + 2].height = 8
        labels = ["Week ending", "Guests", "Spend actual", "Spend teammate", "Spend difference",
                  "Wine actual", "Wine teammate", "Wine difference"]
        for column, label in enumerate(labels, 25):
            _text(ws, current + 3, column, column, label, fill=BURGUNDY,
                  color="FFFFFF", bold=True, align="center", height=36)
        for offset, week in enumerate(weeks, 4):
            row = current + offset
            detail = detail_map.get((*_key(item), week))
            values = [week]
            if detail:
                values += [detail.get("guest_count"), detail.get("check_average"), detail.get("peer_spg"),
                           detail.get("spg_gap"), _wine(detail.get("wine_pct")),
                           _wine(detail.get("peer_wine_pct")), _wine(detail.get("wine_gap"))]
            else:
                values += ["No activity"] + [None] * 6
            for column, value in enumerate(values, 25):
                cell = _text(ws, row, column, column, value if value is not None else "—",
                             fill=PALE if offset % 2 == 0 else "FFFFFF", align="center", height=22)
                cell.number_format = "mm/dd/yyyy" if column == 25 else "#,##0" if column == 26 else (
                    SIGNED_CURRENCY if column in (29, 32) else CURRENCY)
        row = current + 4 + len(weeks)
        _text(ws, row, 25, 32,
              f"Week-to-week steadiness: {item.get('consistency_note') or item.get('consistency', 'Unavailable')}",
              fill=BLUE, height=36)
        _text(ws, row + 1, 25, 32, f"{item.get('coverage_note', '')} A dash in a comparison column means "
              "the metric or eligible teammate reference was unavailable; it does not mean zero.", height=40)
        _text(ws, row + 2, 25, 32, "Each weekly reference is the median of other eligible named people in "
              "the same store and week. Small person samples remain visible. Check shift and guest context "
              "before interpreting a one-week result.", color=MUTED, height=36)
        ws.row_dimensions[row + 3].height = 8
        current = row + 4
    return current - 1


def write_staff_weekly(wb, model: dict[str, Any]) -> None:
    feedback = _feedback(model)
    rows, weeks = list(feedback["rows"]), list(feedback.get("week_ends", []))
    detail_map = {(*_key(item), item["week_end"]): item for item in feedback["detail_rows"]}
    end = max(12, len(weeks) + 6)
    ws = _sheet(wb, "Weekly Performance", "Weekly staff results | actuals, teammate gaps and steadiness", end)
    _text(ws, 3, 1, end, "Read weeks from oldest to newest. Each cell shows actual spend per guest, actual wine "
          "sales per $100, the difference from typical teammates in parentheses, and guests served.",
          fill=BLUE, height=38)
    _text(ws, 4, 1, end, "No recorded activity means no person row was found that week. Unavailable means a "
          "valid metric or teammate reference was missing. A small sample still has visible results. "
          "Print gives one detailed, readable page per person.", fill=PALE, height=38)
    headers = ["Location", "Server", "Data coverage", "Week-to-week steadiness"]
    headers += [_date(week) for week in weeks]
    headers += ["Latest week activity", "Sample context"]
    records = [[item["location"], item["display_name"], item.get("data_coverage"), item.get("consistency"),
                *[_weekly_cell(detail_map.get((*_key(item), week))) for week in weeks],
                _activity(item), _sample(item)] for item in rows]
    _table(ws, "WeeklyPerformanceTable", 6, headers, records, row_height=64)
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = (22 if column <= 4 else 31)
    start = 10 + len(rows)
    last = _weekly_print(ws, rows, detail_map, weeks, start)
    link = _text(ws, 5, 1, end, "Open the printable weekly detail →", color=BURGUNDY, bold=True, height=24)
    link.hyperlink = f"#'Weekly Performance'!Y{start}"
    ws.freeze_panes = "E7"
    _letter_print(ws, f"Y{start}:AF{last}", f"{start}:{start + 1}")


def _dashboard_section(ws, start: int, title: str, note: str, rows: list[dict[str, Any]],
                       kind: str, fill: str) -> int:
    """Paginate complete person rows with repeated section context at native size."""
    current = start
    pending = list(rows)
    first_page = True
    while first_page or pending:
        ws.row_breaks.append(Break(id=current - 1))
        _text(ws, current, 1, 12, title + (" | continued" if not first_page else ""),
              fill=BURGUNDY, color="FFFFFF", bold=True, size=12, height=26)
        _text(ws, current + 1, 1, 12, note, fill=fill, height=38)
        for first, last, label in ((1, 2, "Person / store"), (3, 6, "Actual vs typical teammate (difference)"),
                                   (7, 8, "Compared sample"), (9, 12, "Suggested next step")):
            _text(ws, current + 2, first, last, label, fill=BLUE, bold=True, height=30)
        current += 3
        used = 94
        if not pending:
            _text(ws, current, 1, 12, "No people in this group for the current snapshot.", height=32)
            current += 1
        while pending:
            item = pending[0]
            values = [f"{item['display_name']}\n{item['location']}", _evidence(item),
                      f"{item.get('compared_weeks', 0)} weeks\n{_guests(item.get('compared_guests'))} guests\n"
                      f"{item.get('data_coverage', 'No comparison')}", _prompt(item, kind)]
            height = max(_line_height(value, chars, minimum=64)
                         for value, chars in zip(values, (18, 36, 18, 36)))
            if used + height > 480 and used > 94:
                break
            pending.pop(0)
            for (first, last), value in zip(((1, 2), (3, 6), (7, 8), (9, 12)), values):
                _text(ws, current, first, last, value, fill=fill if current % 2 else PALE,
                      height=height)
            used += height
            current += 1
        current += 1
        first_page = False
    return current


def write_staff_dashboard(wb, model: dict[str, Any], config: dict[str, Any], readiness: Any) -> None:
    feedback = _feedback(model)
    rows = list(feedback["rows"])
    current = [item for item in rows if item.get("present_latest_week")]
    historical = [item for item in rows if not item.get("present_latest_week")]
    coaching = [item for item in feedback["coaching_rows"] if item.get("present_latest_week")]
    recognition = [item for item in feedback["recognition_rows"] if item.get("present_latest_week")]
    in_topics = {_key(item) for item in coaching + recognition}
    checkins = [item for item in current if _key(item) not in in_topics]
    ws = _sheet(wb, "Performance Dashboard", "Staff feedback | where to coach, recognize and check in")
    for column in range(1, 13):
        ws.column_dimensions[get_column_letter(column)].width = 9.5
    latest = feedback.get("latest_complete_week_end") or model.get("latest_complete_week_end")
    _text(ws, 3, 1, 12, f"Through {_date(latest)} | {len(rows)} named people in the selected weeks | "
          f"{len(current)} with activity latest week | {len(historical)} without activity latest week",
          fill=PALE, height=34)
    _text(ws, 4, 1, 12, "Use the lists below for practical conversations. Everyone has a scorecard. "
          "Results describe selling outcomes; ask about shifts and guests before choosing a next step. "
          "Small samples remain visible and are labeled Early indication.", fill=BLUE, height=44)
    _text(ws, 6, 1, 12, "OVERALL STORE RESULTS | latest complete week", fill=BURGUNDY,
          color="FFFFFF", bold=True, size=12, height=26)
    overall = model.get("overall_store") or {}
    baseline = overall.get("baseline_week_count", 0)
    cards = [(1, 3, "Gross sales", overall.get("gross_sales"), "$#,##0", "gross_sales_change_pct", "percent"),
             (4, 5, "Guests", overall.get("guest_count"), "#,##0", "guest_count_change_pct", "percent"),
             (6, 8, "Spend per guest", overall.get("check_average"), CURRENCY, "check_average_change", "dollars"),
             (9, 10, "Wine per $100", _wine(overall.get("wine_pct")), CURRENCY, "wine_pct_change", "wine"),
             (11, 12, "Complete stores", f"{overall.get('complete_location_count', 0)} of {len(config.get('locations', {}))}",
              "General", None, None)]
    for first, last, label, value, fmt, change_key, kind in cards:
        _text(ws, 7, first, last, label, fill=BLUE, bold=True, align="center", height=30)
        cell = _text(ws, 8, first, last, value if value is not None else "Unavailable",
                     fill=BLUE, bold=True, size=18, align="center", height=38)
        cell.number_format = fmt
        delta = overall.get(change_key) if change_key else None
        if change_key is None:
            note = "Source store totals"
        elif delta is None or not baseline:
            note = "No prior complete-week baseline"
        else:
            difference = f"{delta:+.1%}" if kind == "percent" else _money(delta * 100 if kind == "wine" else delta, signed=True)
            note = f"Change: {difference} vs prior {baseline}-week baseline"
        _text(ws, 9, first, last, note, fill=BLUE, color=MUTED, align="center", height=44)
    _text(ws, 11, 1, 12, f"{len(coaching)} coaching conversations | {len(recognition)} recognition opportunities | "
          f"{len(checkins)} routine or near-threshold check-ins", fill=GREEN, bold=True, size=12, height=34)
    _text(ws, 12, 1, 12, "A person can appear in both coaching and recognition when the two metrics point "
          "in different directions. The lists prioritize people with recorded activity in the latest week. "
          "Historical results are shown separately below.", fill=PALE, height=44)
    _text(ws, 13, 1, 12, "Read results as actual vs typical teammate, followed by the difference. Spend is dollars "
          "per guest. Wine is dollars of wine in each $100 of sales. Samples and steadiness travel with every result.",
          color=MUTED, height=40)
    row = 16
    sections = [
        ("COACHING CONVERSATIONS", "Lower selling results suggest a topic to explore. Use an early indication "
         "as a question and verify shift context before setting a goal.", coaching, "coaching", AMBER),
        ("RECOGNITION OPPORTUNITIES", "Recognize the specific result and ask which habits helped. "
         "Share useful habits with the team after checking the shift context.", recognition, "recognition", GREEN),
        ("ROUTINE AND NEAR-THRESHOLD CHECK-INS", "Keep ordinary check-ins going. Near-threshold results "
         "do not receive a firm positive or negative label; collect more shifts and revisit.", checkins, "checkin", BLUE),
        ("NO RECORDED ACTIVITY LATEST WEEK", "These people were recorded earlier in the selected window. "
         "Absence from the latest source week does not establish employment status. Confirm context before acting.",
         historical, "historical", PALE),
    ]
    for title, note, people, kind, fill in sections:
        row = _dashboard_section(ws, row, title, note, people, kind, fill)
    _letter_print(ws, f"A1:L{row - 1}", "1:2")
    ws.freeze_panes = "A5"


def write_staff_methodology(wb, model: dict[str, Any], config: dict[str, Any]) -> None:
    feedback = _feedback(model)
    thresholds = feedback.get("thresholds") or {}
    minimum_peers = thresholds.get("minimum_peers", 5)
    score_limits = config.get("management_score_thresholds", {})
    spg_limits = {"neutral": 11.5, "strong": 17.5, **score_limits.get("check_average", {})}
    wine_limits = {"neutral": .041, "strong": .057, **score_limits.get("wine_pct", {})}
    ws = _sheet(wb, "Methodology", "How to read the staff feedback")
    for column in range(1, 13):
        ws.column_dimensions[get_column_letter(column)].width = 9.5
    _text(ws, 3, 1, 12, "A practical guide for managers and staff. The same definitions apply to the "
          f"dashboard, scorecards and weekly results. Methodology: {model.get('methodology_version', '2026.10-v3')}.",
          fill=BLUE, height=34)
    sections = [
        ("1. Two selling results", "Spend per guest is gross sales divided by guests. Wine sales per $100 is "
         "wine sales divided by gross sales, multiplied by 100. For example, $12 in wine per $100 of sales "
         "compared with $10 for typical teammates is a $2 difference. Higher or lower does not by itself "
         "describe service quality or total job performance."),
        ("2. A fairer teammate reference", f"Each person is compared with the median of other eligible named "
         f"people in the same store and week; at least {minimum_peers} other eligible people are required. "
         "The median is the middle teammate result. The subject is excluded from their own reference. "
         "The existing peer eligibility rules still apply to reference members. A person does not need "
         "to meet that larger sample gate to see their own descriptive comparison."),
        ("3. Every valid person sample stays visible", "A named person with positive sales, positive guests "
         "and recorded activity can have descriptive results even with one guest or one week. Small samples "
         "are explicitly labeled; no made-up comparison is shown when a metric or teammate reference is "
         "missing. Shared, area and unresolved identities remain outside individual staff feedback."),
        ("4. Reading the eight-week result", "The window uses up to eight complete shared weeks. A week's "
         "comparison receives weight equal to that person's guests, capped at 50. This keeps one busy week "
         "from dominating. Actual and typical teammate values use the same weeks and weights for each "
         "metric, so their difference is the displayed gap. All-recorded-work totals can include additional "
         "weeks without a valid reference and need not equal the comparison values."),
        ("5. How much evidence is behind it?", "Established means at least 6 compared weeks and 200 guests. "
         "Building means at least 4 compared weeks and 150 guests. Other valid comparisons are Early "
         "indication. No comparison means the needed denominator or teammate reference is missing. "
         "Weeks and guests are shown beside results so people can judge the strength of the evidence. "
         "These labels describe data coverage, not the employee."),
        ("6. Coaching, recognition and close calls", "A difference of $5 per guest or $1.50 in wine per $100 "
         "starts the conversation screen. Lower results suggest coaching; higher results suggest "
         "recognition. Values within 10% of either threshold are labeled near the conversation threshold "
         "and do not receive a firm positive or negative category for that metric. Display precision is "
         "used consistently at the boundary. A person can appear in both lists when metrics point in "
         "different directions. Small samples always remain tentative."),
        ("7. Week-to-week steadiness", "Steady, Some variation and Variable describe how much the person's "
         "teammate gaps change from week to week, separately from whether the results are higher or "
         "lower. One week cannot establish a pattern. The technical measure is the sample standard "
         f"deviation of weekly gaps: Steady is at most ${spg_limits['neutral']:.2f} for spend and "
         f"${wine_limits['neutral'] * 100:.2f} for wine per $100; Some variation is at most "
         f"${spg_limits['strong']:.2f} and ${wine_limits['strong'] * 100:.2f}; larger movement is Variable. "
         "Read the weekly table and sample "
         "note together; steadiness is not the same as statistical certainty."),
        ("8. Make the conversation useful", "Ask which shifts, sections, parties, promotions and guest "
         "preferences shaped the result. For coaching, practice one relevant habit and agree on a "
         "follow-up after more shifts. For recognition, name the specific result and invite the person "
         "to share what worked. For an early or near-threshold indication, ask first and collect more "
         "evidence. No recorded activity latest week is a source-data statement, not an employment-status "
         "determination."),
        ("9. Scope and printing", "Store headline totals come from store source totals, not summed server "
         "rows. The staff conversation lists are descriptive and separate from the qualified Management "
         "Center Current Actions workflow. Print the dashboard for team review. Print Server Scorecards or Weekly "
         "Performance for one native Letter landscape page per person. Their compact print views are "
         "linked from the top of each sheet. Generated results are protected; use the authorized management "
         "inputs for action notes."),
    ]
    row, used = 5, 90
    for heading, text in sections:
        height = _line_height(text, 112, minimum=50)
        if used + 28 + height > 500:
            ws.row_breaks.append(Break(id=row - 1))
            used = 56
        _text(ws, row, 1, 12, heading, fill=BURGUNDY, color="FFFFFF", bold=True, size=11, height=28)
        _text(ws, row + 1, 1, 12, text, fill=PALE, height=height)
        used += 28 + height + 12
        row += 3
        ws.row_dimensions[row - 1].height = 12
    _letter_print(ws, f"A1:L{row - 1}", "1:2")
    ws.freeze_panes = "A5"
